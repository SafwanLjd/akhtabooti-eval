import json
import random
import shutil
from collections import Counter
from collections.abc import Mapping, Sequence
from dataclasses import replace
from pathlib import Path
from typing import Any

import docx
import pytest
from docx.oxml.ns import qn
from PIL import Image

from akheval.dataset import DOCUMENTS, LANGUAGES, Dataset, Rendering, validate
from akheval.synth.build import (
    DocSpec,
    build_dataset,
    generate_document,
    generate_documents,
    limit_documents,
    plan_documents,
    render_images,
)
from akheval.synth.builder import block_side, has_back
from akheval.synth.config import load_config, parse_config
from akheval.synth.doctypes import CARD_IMAGES, CHAT_IMAGES, DOC_TYPES, NO_IMAGES, PAGE_IMAGES
from akheval.synth.imaging import plan_images
from akheval.synth.native import DIRECT_WRITERS, native_relpath
from akheval.synth.readback import FileText, extractous_readback, missing_values

CONFIG = Path(__file__).resolve().parents[1] / "configs" / "dataset-v1.toml"
STRATA = len(DOC_TYPES) * len(LANGUAGES)


def stub_renderer(
    layout: Sequence[Mapping[str, Any]],
    *,
    style: str,
    language: str,
    capture: str,
    severity: str | None,
    rng: random.Random,
    out: Path,
) -> dict[str, Any]:
    """A blank image: enough for the build to record and checksum a rendering."""
    del language, capture, severity
    out.parent.mkdir(parents=True, exist_ok=True)
    Image.new("L", (32, 32), 255).save(out, "PDF" if out.suffix == ".pdf" else None)
    sides = sorted({block_side(b) for b in layout})
    return {"stub": True, "style": style, "blocks": len(layout), "sides": sides, "draw": round(rng.random(), 6)}


def test_repository_config_names_every_document_type() -> None:
    cfg = load_config(CONFIG)
    assert set(cfg.documents) == set(DOC_TYPES)
    assert cfg.dev_fraction == pytest.approx(0.3)
    assert sum(cfg.images.severity.values()) == pytest.approx(1.0)


def test_config_rejects_unknown_document_types() -> None:
    raw = dict(load_config(CONFIG).raw)
    raw["documents"] = {"letter": 1, "fax": 2}
    with pytest.raises(ValueError, match="fax"):
        parse_config(raw)


def test_limit_takes_one_document_per_stratum_first() -> None:
    specs = plan_documents(load_config(CONFIG).documents)
    chosen = limit_documents(specs, STRATA)
    assert len(chosen) == STRATA
    assert {(s.doc_type, s.language) for s in chosen} == {(t, lang) for t in DOC_TYPES for lang in LANGUAGES}
    assert [s.id for s in chosen] == sorted(s.id for s in chosen)


def test_native_formats_rotate_through_every_format_of_a_type() -> None:
    for name, doc_type in DOC_TYPES.items():
        n = len(doc_type.native_formats)
        for lang in LANGUAGES:
            formats = {DocSpec(name, lang, i).native_format for i in range(1, n + 1)}
            assert formats == set(doc_type.native_formats)


def test_generation_is_deterministic() -> None:
    cfg = load_config(CONFIG)
    first = [json.dumps(d.to_dict(), ensure_ascii=False) for d in generate_documents(cfg, STRATA)]
    second = [json.dumps(d.to_dict(), ensure_ascii=False) for d in generate_documents(cfg, STRATA)]
    assert first == second


def test_more_documents_leave_existing_documents_unchanged() -> None:
    cfg = load_config(CONFIG)
    bigger = replace(cfg, documents={k: v + 3 for k, v in cfg.documents.items()})
    small = {d.id: d for d in generate_documents(cfg)}
    large = {d.id: d for d in generate_documents(bigger)}
    assert set(small) < set(large)
    for doc_id, doc in small.items():
        assert large[doc_id].text == doc.text
        assert large[doc_id].spans == doc.spans


def test_split_is_stratified_by_type_and_language() -> None:
    cfg = load_config(CONFIG)
    docs = generate_documents(cfg)
    for (doc_type, lang), count in Counter((d.doc_type, d.language) for d in docs).items():
        dev = sum(1 for d in docs if d.doc_type == doc_type and d.language == lang and d.split == "dev")
        assert dev == round(count * cfg.dev_fraction)


def test_image_plans_follow_the_type_policy() -> None:
    shares = load_config(CONFIG).images
    rng = random.Random(0)
    cards = [plan_images(CARD_IMAGES, shares, rng) for _ in range(400)]
    assert all(jobs[0].capture == "photo" and jobs[0].fmt == "jpg" for jobs in cards)
    assert 0.2 < sum(len(jobs) == 2 for jobs in cards) / len(cards) < 0.4
    pages = [plan_images(PAGE_IMAGES, shares, rng) for _ in range(400)]
    assert 0.4 < sum(bool(jobs) for jobs in pages) / len(pages) < 0.6
    chats = [job for _ in range(400) for job in plan_images(CHAT_IMAGES, shares, rng)]
    assert all(job.capture == "digital" and job.fmt == "png" and job.severity is None for job in chats)
    assert plan_images(NO_IMAGES, shares, rng) == []
    severities = Counter(job.severity for jobs in cards for job in jobs)
    assert severities["harsh"] < severities["mild"]


def test_two_sided_cards_plan_one_photo_per_side() -> None:
    shares = load_config(CONFIG).images
    plans = [plan_images(CARD_IMAGES, shares, random.Random(seed), two_sided=True) for seed in range(40)]
    for jobs in plans:
        front, back, *scans = jobs
        assert (front.face, back.face) == ("front", "back")
        assert front.severity == back.severity
        assert front.rendering_id("d") == "d.photo-front.jpg"
        assert back.relpath("d") == "files/d/photo-back.jpg"
        assert all(job.capture == "scan" and job.face is None for job in scans)
    assert any(len(jobs) == 3 for jobs in plans)


def test_photos_of_two_sided_cards_show_one_side_each(tmp_path: Path) -> None:
    cfg = load_config(CONFIG)
    documents = generate_documents(cfg, STRATA)
    images = render_images(documents, tmp_path, cfg, stub_renderer)
    two_sided = [d for d in documents if has_back(d.meta["layout"])]
    assert {d.doc_type for d in two_sided} >= {"jo_id_card", "sy_id_card"}
    for doc in documents:
        doc.renderings = tuple(images[doc.id])
        if doc not in two_sided:
            assert all(r.visible is None for r in doc.renderings), doc.id
            continue
        photos = [r for r in doc.renderings if r.capture == "photo"]
        assert [r.path for r in photos] == [f"files/{doc.id}/photo-front.jpg", f"files/{doc.id}/photo-back.jpg"]
        assert [r.params["sides"] for r in photos] == [["front"], ["back"]]
        shown = sorted((s for r in photos for s in doc.visible_spans(r)), key=lambda s: (s.start, s.end))
        assert shown == list(doc.spans), doc.id
        assert all(r.visible is None for r in doc.renderings if r.capture == "scan")
    assert validate(Dataset(root=tmp_path, manifest={}, documents=documents)) == []


@pytest.mark.parametrize("fmt", sorted(DIRECT_WRITERS))
def test_direct_native_files_contain_every_gold_value(tmp_path: Path, fmt: str) -> None:
    for doc in generate_documents(load_config(CONFIG), STRATA):
        (tmp_path / "files" / doc.id).mkdir(parents=True)
        DIRECT_WRITERS[fmt](doc, tmp_path / native_relpath(doc.id, fmt))
        doc.renderings = (Rendering(f"{doc.id}.digital.{fmt}", native_relpath(doc.id, fmt), fmt, "native", "digital"),)
        assert missing_values(doc, tmp_path) == []


def test_docx_tables_size_arabic_text_like_latin_text(tmp_path: Path) -> None:
    doc = generate_document(DocSpec("contact_list", "ar", 3), seed=load_config(CONFIG).seed, split="test")
    path = tmp_path / "list.docx"
    DIRECT_WRITERS["docx"](doc, path)
    runs = [r for cell in docx.Document(str(path)).tables[0]._cells for p in cell.paragraphs for r in p.runs]
    sizes = {(r._r.rPr.find(qn("w:sz")).get(qn("w:val")), r._r.rPr.find(qn("w:szCs")).get(qn("w:val"))) for r in runs}
    assert sizes == {("16", "16")}


@pytest.mark.parametrize(
    ("text", "fmt", "value", "found"),
    [
        ("الهاتف 4567 123 79 962+", "pdf", "+962 79 123 4567", True),
        ("الهاتف 4567 123 79 962+", "txt", "+962 79 123 4567", False),
        ("الرقم AB\n123 نهاية", "pdf", "AB-123", True),
        ("962\nx\ny\nz\n79 123 4567", "pdf", "+962 79 123 4567", False),
    ],
)
def test_file_text_tolerates_pdf_text_layers_only(text: str, fmt: str, value: str, found: bool) -> None:
    assert FileText(text, fmt).contains(value) is found


requires_office = pytest.mark.skipif(
    shutil.which("soffice") is None or shutil.which("pdftotext") is None, reason="needs LibreOffice and poppler"
)


@pytest.mark.integration
@requires_office
def test_build_writes_a_valid_reproducible_dataset(tmp_path: Path) -> None:
    first = build_dataset(CONFIG, tmp_path / "a", limit=STRATA, renderer=stub_renderer)
    assert validate(first) == []
    formats = Counter(r.format for _, r in first.renderings() if r.channel == "native")
    assert formats["pdf"] and formats["odt"]
    assert first.manifest["generator"]["config"]["dataset"]["seed"] == load_config(CONFIG).seed
    recovery = first.manifest["generator"]["extractous_recovery"]
    assert set(recovery) == set(formats)
    assert all(0 < cell["verbatim"] <= cell["values"] for by_lang in recovery.values() for cell in by_lang.values())
    build_dataset(CONFIG, tmp_path / "b", limit=STRATA, renderer=stub_renderer)
    assert (tmp_path / "a" / DOCUMENTS).read_bytes() == (tmp_path / "b" / DOCUMENTS).read_bytes()
    reloaded = Dataset.load(tmp_path / "a")
    assert [d.id for d in reloaded.documents] == [d.id for d in first.documents]


@pytest.mark.integration
def test_extractous_readback_reports_files_it_cannot_parse(tmp_path: Path) -> None:
    pytest.importorskip("extractous")
    documents = generate_documents(load_config(CONFIG), 2)
    for doc, fmt in zip(documents, ("txt", "docx"), strict=True):
        (tmp_path / "files" / doc.id).mkdir(parents=True)
        DIRECT_WRITERS[fmt](doc, tmp_path / native_relpath(doc.id, fmt))
        doc.renderings = (Rendering(f"{doc.id}.digital.{fmt}", native_relpath(doc.id, fmt), fmt, "native", "digital"),)
    (tmp_path / native_relpath(documents[1].id, "docx")).write_bytes(b"PK\x03\x04 truncated archive")
    recovery, failures = extractous_readback(documents, tmp_path)
    assert len(failures) == 1
    assert documents[1].id in failures[0]
    cell = recovery["txt"][documents[0].language]
    assert cell["verbatim"] == cell["values"] == len(documents[0].spans)


@pytest.mark.integration
@requires_office
def test_build_refuses_to_overwrite_without_permission(tmp_path: Path) -> None:
    build_dataset(CONFIG, tmp_path, limit=1, renderer=stub_renderer)
    with pytest.raises(FileExistsError):
        build_dataset(CONFIG, tmp_path, limit=1, renderer=stub_renderer)
    manifest = json.loads((tmp_path / "manifest.json").read_text(encoding="utf-8"))
    (tmp_path / "manifest.json").write_text(json.dumps({**manifest, "frozen": True}), encoding="utf-8")
    with pytest.raises(FileExistsError, match="frozen"):
        build_dataset(CONFIG, tmp_path, limit=1, renderer=stub_renderer, overwrite=True)


@pytest.mark.integration
@requires_office
def test_build_with_the_project_image_renderer(tmp_path: Path) -> None:
    images = pytest.importorskip("akheval.synth.images")
    dataset = build_dataset(CONFIG, tmp_path, limit=STRATA, renderer=images.render_image)
    assert validate(dataset) == []
    assert any(r.channel == "ocr" for _, r in dataset.renderings())
