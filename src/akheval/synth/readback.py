"""Read native renderings back to text: check that every gold value survived and that extractous parses each file."""

from __future__ import annotations

import json
import re
import subprocess
import unicodedata
import zipfile
from collections import Counter
from collections.abc import Callable, Iterator, Sequence
from html.parser import HTMLParser
from pathlib import Path
from xml.etree import ElementTree

import docx
from openpyxl import load_workbook

from ..dataset import Document, Rendering
from ..matching import BIDI_CONTROLS

ARABIC_MARKS = re.compile("[\u064b-\u0652\u0670\u0640]")
WHITESPACE = re.compile(r"\s+")
WORD = re.compile(r"\w+")
SIGNS = re.compile(r"[-+]")
PDFTOTEXT_TIMEOUT = 120
PDF_LINE_WINDOW = 3
"""Consecutive pdftotext lines searched together, so a value that wrapped onto the next line still matches."""

Recovery = dict[str, dict[str, dict[str, int]]]
"""Per native format and language: gold values in the files, and how many of them extractous returns verbatim."""


class _TextCollector(HTMLParser):
    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.parts: list[str] = []

    def handle_data(self, data: str) -> None:
        self.parts.append(data)


def _read_html(path: Path) -> str:
    parser = _TextCollector()
    parser.feed(path.read_text(encoding="utf-8"))
    return "\n".join(parser.parts)


def _read_json(path: Path) -> str:
    return "\n".join(json.loads(path.read_text(encoding="utf-8"))["lines"])


def _read_docx(path: Path) -> str:
    document = docx.Document(str(path))
    lines = [p.text for p in document.paragraphs]
    lines += [cell.text for table in document.tables for row in table.rows for cell in row.cells]
    return "\n".join(lines)


def _read_xlsx(path: Path) -> str:
    workbook = load_workbook(path, read_only=True)
    lines = []
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows(values_only=True):
            lines.append(" ".join(str(v) for v in row if v is not None))
    workbook.close()
    return "\n".join(lines)


def _read_odt(path: Path) -> str:
    with zipfile.ZipFile(path) as archive:
        root = ElementTree.fromstring(archive.read("content.xml"))
    paragraphs = (e for e in root.iter() if e.tag.endswith(("}p", "}h")))
    return "\n".join("".join(e.itertext()) for e in paragraphs)


def _read_pdf(path: Path) -> str:
    result = subprocess.run(
        ["pdftotext", "-enc", "UTF-8", str(path), "-"], capture_output=True, check=True, timeout=PDFTOTEXT_TIMEOUT
    )
    return result.stdout.decode("utf-8")


def _read_plain(path: Path) -> str:
    return path.read_text(encoding="utf-8")


READERS: dict[str, Callable[[Path], str]] = {
    "txt": _read_plain,
    "csv": _read_plain,
    "html": _read_html,
    "json": _read_json,
    "docx": _read_docx,
    "xlsx": _read_xlsx,
    "odt": _read_odt,
    "pdf": _read_pdf,
}


def read_native_text(path: Path, fmt: str) -> str:
    return READERS[fmt](path)


def _clean(text: str) -> str:
    """NFKC, bidi controls turned into spaces, Arabic diacritics and tatweel removed."""
    return ARABIC_MARKS.sub("", BIDI_CONTROLS.sub(" ", unicodedata.normalize("NFKC", text)))


def normalise(text: str) -> str:
    """Cleaned text without whitespace."""
    return WHITESPACE.sub("", _clean(text))


def _words(text: str) -> Counter[str]:
    return Counter(WORD.findall(_clean(text)))


def _in_visual_order(value: str, lines: list[str]) -> bool:
    """True when the words of ``value`` appear within a few consecutive lines, in any order.

    pdftotext writes digit groups inside right-to-left lines in visual order, so
    "+962 79 123 4567" can come out as "4567 123 79 962+".
    """
    needed = _words(value)
    for start in range(len(lines)):
        window = _words(" ".join(lines[start : start + PDF_LINE_WINDOW]))
        if not needed - window:
            return True
    return False


def _found_in_pdf_text(value: str, flat: str, lines: list[str]) -> bool:
    """Accept PDF text that lost a hyphen at a line break, moved a plus sign or printed digit groups in visual order."""
    signless = SIGNS.sub("", normalise(value))
    return signless in SIGNS.sub("", flat) or _in_visual_order(value, lines)


class FileText:
    """Text read from one file, normalised once for value lookups."""

    def __init__(self, text: str, fmt: str) -> None:
        self.fmt = fmt
        self.flat = normalise(text)
        self.lines = text.splitlines()

    def verbatim(self, value: str) -> bool:
        return normalise(value) in self.flat

    def contains(self, value: str) -> bool:
        """Verbatim, or in a PDF, with the hyphen losses and visual-order digit groups of PDF text layers."""
        return self.verbatim(value) or (self.fmt == "pdf" and _found_in_pdf_text(value, self.flat, self.lines))


def _natives(doc: Document) -> Iterator[Rendering]:
    return (r for r in doc.renderings if r.channel == "native")


def missing_values(doc: Document, root: Path) -> list[str]:
    """Gold values that the native renderings of ``doc`` do not contain."""
    problems: list[str] = []
    for rendering in _natives(doc):
        text = FileText(read_native_text(root / rendering.path, rendering.format), rendering.format)
        problems += [
            f"{doc.id}: {rendering.path} lacks {v.label} {v.value!r}"
            for v in doc.visible_values(rendering)
            if not text.contains(v.value)
        ]
    return problems


def extractous_readback(documents: Sequence[Document], root: Path) -> tuple[Recovery, list[str]]:
    """Read every native file with extractous, the library the desktop app uses.

    Returns the gold values extractous recovers verbatim per format and language, and one
    problem per file it cannot parse: desktop release v0.1.0 stops a folder scan at such a file.
    """
    # Imported here so missing_values runs without extractous; the synth and extract extras install it.
    from ..extractors.base import ExtractionError
    from ..extractors.extractous import ExtractousExtractor

    extractor = ExtractousExtractor()
    recovery: Recovery = {}
    failures: list[str] = []
    for doc in documents:
        for rendering in _natives(doc):
            try:
                text = FileText(extractor.extract(root / rendering.path), rendering.format)
            except ExtractionError as exc:
                failures.append(f"{doc.id}: {exc}")
                continue
            values = doc.visible_values(rendering)
            cell = recovery.setdefault(rendering.format, {}).setdefault(doc.language, {"values": 0, "verbatim": 0})
            cell["values"] += len(values)
            cell["verbatim"] += sum(text.verbatim(v.value) for v in values)
    return {fmt: dict(sorted(by_lang.items())) for fmt, by_lang in sorted(recovery.items())}, failures
